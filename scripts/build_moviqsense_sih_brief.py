from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docs" / "MovIqSense_SIH_2026_Project_Brief.docx"

NAVY = "132238"
BLUE = "1F4E79"
PALE_BLUE = "EAF2F8"
LIGHT_BLUE = "F5F9FC"
GRAY = "D9E2F3"
TEXT = "1F2937"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_text(paragraph, text, bold=False, color=TEXT, size=10.5, italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text, size=10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    add_text(p, text, size=10.2)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = "Aptos Display"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.size = Pt(15 if level == 1 else 12)
    run.bold = True
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, item in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, item, bold=True, color="FFFFFF", size=9.2)
    for row_no, row in enumerate(rows):
        cells = table.add_row().cells
        for i, item in enumerate(row):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_shading(cell, "FFFFFF" if row_no % 2 == 0 else LIGHT_BLUE)
            set_cell_border(cell)
            set_cell_margin(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, item, size=9.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def make_document():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp, "MovIqSense | Smart India Hackathon 2026 | Problem Statement 26124", color="667085", size=8.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    r = title.add_run("MovIqSense")
    r.font.name = "Aptos Display"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    r.font.size = Pt(27)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(14)
    add_text(sub, "AI Powered Mobile Urban Intelligence Platform Using Public Transport Fleet", bold=True, color=BLUE, size=13)

    add_table(doc, ["Submission Detail", "Value"], [
        ["Problem Statement ID", "26124"],
        ["Organisation and Department", "Bharat Electronics Limited"],
        ["Theme", "Smart Automation"],
        ["Category", "Software"],
        ["Target Domain", "GovTech, intelligent urban mobility, road safety, civic maintenance and public transport analytics"],
        ["Implementation Scope", "Pilot corridor to city fleet, State Transport Undertaking and national Smart Mobility scale"],
    ], [2.05, 5.65])

    add_heading(doc, "Executive Summary")
    add_body(doc, "MovIqSense transforms buses equipped with cameras and GPS into mobile urban sensing units. It applies artificial intelligence to bus camera feeds to identify road defects, waterlogging, missing road infrastructure, traffic bottlenecks, vulnerable pedestrian situations, accident evidence and vehicle plate crops. Each verified event is associated with bus identity, time, GPS location, confidence and visual evidence, then displayed on a central government GIS dashboard.")
    add_body(doc, "The solution addresses a practical gap in urban governance: bus cameras generally record incidents but do not create real-time civic intelligence. MovIqSense converts existing fleet movement into a cooperative sensing network. It enables authorities to see high-risk defects, traffic conditions and safety events sooner, prioritise maintenance work, coordinate response and monitor closure through an evidence-based workflow.")

    add_heading(doc, "The Problem and Public Value")
    add_body(doc, "Public buses travel across major roads, interior streets, commercial areas, school zones and flood-prone corridors every day. Yet road-condition monitoring still relies heavily on fixed CCTV, periodic manual inspections and citizen complaints. These methods leave coverage gaps, delay verification and make it difficult to rank which road defect or incident requires urgent action.")
    add_body(doc, "The Ministry of Road Transport and Highways recorded 4,80,583 road accidents, 1,72,890 fatalities and 4,62,825 injuries in 2023. Road hazards, unsafe crossings, congestion and delayed evidence collection create avoidable public-safety and operational risks. MovIqSense supports the national direction of Intelligent Transport Systems, Digital India, Smart Cities and technology-enabled public transport.")

    add_heading(doc, "Proposed Solution")
    add_body(doc, "MovIqSense is an edge-ready AI and GIS platform with two connected layers. The field layer captures bus-camera video and GPS information. The central layer aggregates geo-tagged events from the fleet, removes duplicates, ranks priorities and presents actionable information to transport authorities, road-maintenance teams and traffic police.")
    add_table(doc, ["Field Layer", "Central Command Layer"], [
        ["Bus camera or uploaded video; bus number; route data; GPS latitude, longitude and timestamp", "GIS map, live fleet view, defect register, traffic analytics, accident alerts, evidence history and repair workflow"],
        ["AI detects road hazards, vehicles, pedestrians, accidents and plate evidence", "Multi-bus verification clusters repeated observations and reduces duplicate event creation"],
        ["Sends compact structured JSON events and selected evidence rather than continuous raw video", "Ranks incidents by severity, confidence, persistence, traffic impact and safety exposure"],
    ], [3.85, 3.85])

    add_heading(doc, "Core Capabilities")
    capabilities = [
        ("Road and civic infrastructure monitoring", "Detects potholes, damaged road surfaces, waterlogging, speedbreakers, zebra crossings, road dividers, traffic signs and other road hazards."),
        ("Traffic intelligence", "Counts cars, buses, motorcycles and pedestrians; classifies traffic density and identifies possible bottlenecks and route-delay signals."),
        ("Pedestrian safety", "Flags pedestrian proximity to active traffic, with focus on school zones, bus stops and unsafe crossings."),
        ("Accident and hit-and-run support", "Creates GPS-tagged accident evidence records with timestamp, severity, vehicle crop and plate-crop evidence for authorised review."),
        ("Government action workflow", "Shows defects and incidents on a GIS map; supports status updates, repair deadlines, AI priority queues and official incident PDF reports."),
    ]
    for head, text in capabilities:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_text(p, head + ": ", bold=True, color=NAVY, size=10.4)
        add_text(p, text, size=10.4)

    add_heading(doc, "Operational Workflow")
    add_table(doc, ["Step", "MovIqSense Process", "Government Outcome"], [
        ["1. Capture", "A bus starts a route with camera and GPS enabled.", "The vehicle becomes an active mobile sensing point."],
        ["2. Detect", "AI analyses selected frames for road hazards, traffic and incident evidence.", "Detection confidence and visual evidence are generated."],
        ["3. Localise", "Bus ID, timestamp, latitude and longitude are attached to the event.", "The event becomes visible on the GIS map."],
        ["4. Verify", "Repeated observations across frames and buses are clustered.", "False or duplicate alerts are reduced."],
        ["5. Prioritise", "Risk score considers severity, confidence, persistence and traffic impact.", "High-risk defects and incidents move to the top of the queue."],
        ["6. Act and close", "Officer assigns response, repair deadline and status.", "The authority monitors Detected, In Progress and Solved closure states."],
    ], [0.65, 3.6, 3.45])

    add_heading(doc, "Technology Architecture")
    add_table(doc, ["Layer", "Technology and Role"], [
        ["Computer vision", "Ultralytics YOLO and PyTorch for vehicle, pedestrian, pothole, waterlogging, road-hazard, accident and plate-crop detection."],
        ["Video processing", "OpenCV and browser Canvas APIs for frame extraction, resizing, annotation and downloadable processed MP4 evidence."],
        ["Tracking and analytics", "Multi-frame verification, ByteTrack-compatible association, Kalman-filter-ready smoothing, object counting and density classification."],
        ["Backend", "Python, FastAPI and REST APIs for frame inference, location ingestion, events, buses, defects and incident reports."],
        ["Data and GIS", "SQLite for prototype persistence; Leaflet and OpenStreetMap for live locations, hazards, incidents and priority-zone mapping."],
        ["Production path", "Edge device inference, event-based JSON telemetry, secure object storage, PostgreSQL/PostGIS, role-based access and government API integration."],
    ], [1.65, 6.05])

    add_heading(doc, "Business Model Canvas")
    canvas = [
        ("Customer segments", "State Transport Undertakings, Smart City ICCCs, municipal corporations, PWD, traffic police, fleet operators and citizens."),
        ("Value proposition", "Turns existing bus cameras into mobile AI sensors; provides earlier geo-tagged defect discovery, traffic intelligence, incident evidence and prioritised government action."),
        ("Channels", "SIH pilot, STU proof of concept, Smart City partnerships, system integrators, government tenders and ICCC integration."),
        ("Customer relationships", "Dedicated pilot support, officer training, dashboard configuration, alert review workflow, annual maintenance and model improvement."),
        ("Revenue streams", "Per-bus or fleet software licence, implementation and integration services, annual maintenance, analytics reports and staff training."),
        ("Key activities", "AI model validation, API and dashboard development, fleet integration, GIS analytics, security controls, support and pilot measurement."),
        ("Key resources", "AI models, bus camera and GPS feeds, secure backend, maps, validated datasets, government workflow knowledge and integration partnerships."),
        ("Key partnerships", "BEL, STUs, municipal bodies, PWD, traffic police, Smart City command centres, edge-hardware providers and system integrators."),
        ("Cost structure", "AI/ML development, backend/GIS, edge hardware, secure storage, connectivity, integration, deployment, training and support."),
    ]
    add_table(doc, ["Canvas Block", "MovIqSense Strategy"], canvas, [1.7, 6.0])

    add_heading(doc, "Implementation and Scalability")
    add_body(doc, "The hackathon prototype focuses on a complete working chain: driver video and GPS input, AI detection, geo-tagged event storage, a government dashboard, traffic indicators, road-defect workflow, accident evidence and processed-video download. This proves feasibility without requiring a full fleet deployment during the hackathon.")
    add_table(doc, ["Phase", "Scope", "Expected Result"], [
        ["Hackathon prototype", "One or more video sources, browser GPS, AI model inference, dashboard and SQLite records.", "Demonstrates end-to-end detection-to-action workflow."],
        ["Pilot", "5-10 buses on selected corridors with officer verification.", "Measures alert precision, latency, repair acknowledgement and user adoption."],
        ["City deployment", "100+ buses, integrated GPS/AVL, central command and maintenance systems.", "Creates road-condition maps, congestion patterns and fleet intelligence."],
        ["State and national scale", "STU-wide deployment with interoperable government APIs and PostGIS analytics.", "Builds a reusable smart-mobility intelligence framework."],
    ], [1.35, 3.35, 3.0])

    add_heading(doc, "Expected Impact and Success Metrics")
    add_table(doc, ["Area", "Pilot Target", "Public Value"], [
        ["Road-defect discovery", "Geo-tagged observation generated within minutes of verified detection.", "Faster inspection and preventive repair."],
        ["Incident response", "Accident evidence package with location, time, severity and vehicle crop available immediately after detection.", "Better initial situational awareness for authorised responders."],
        ["Maintenance planning", "Priority ranking based on confidence, recurrence, traffic impact and safety exposure.", "Resources focus first on high-risk defects."],
        ["Coverage", "Existing bus routes provide daily observations across many corridors.", "Broader coverage without creating an entirely new camera network."],
        ["Sustainability", "Event-based telemetry rather than continuous raw-video upload.", "Lower bandwidth use and more practical large-fleet deployment."],
    ], [1.55, 3.4, 2.75])

    add_heading(doc, "Privacy and Responsible Deployment")
    add_body(doc, "MovIqSense is designed as a decision-support platform, not an autonomous enforcement system. AI alerts require authorised human verification before legal action. The production roadmap includes HTTPS transport security, role-based access, audit logs, configurable retention, event-only transmission, controlled access to plate evidence and automatic face or non-relevant plate blurring where policy requires.")

    add_heading(doc, "Why MovIqSense Merits Selection")
    add_body(doc, "MovIqSense has strong problem fit because it directly addresses every major requirement of Problem Statement 26124: mobile bus-based sensing, road-defect detection, traffic analysis, pedestrian safety, incident evidence, GPS localisation, GIS visualisation, route insights and low-bandwidth edge processing. Its innovation lies in converting existing buses into a cooperative, citywide sensing network rather than creating another isolated camera system.")
    add_body(doc, "The proposal is technically feasible, operationally relevant and scalable. It combines a demonstrable software prototype with a realistic path from one bus to a fleet, city, state and national deployment. Most importantly, it gives authorities an actionable answer to a daily urban challenge: see the problem early, know its location, understand its risk, and act before it becomes more dangerous or expensive.")

    add_heading(doc, "References")
    refs = [
        "Ministry of Road Transport and Highways, Road Accidents in India 2023, Government of India.",
        "Ministry of Housing and Urban Affairs, National Urban Transport Policy review, Government of India.",
        "PM-eBus Sewa, Ministry of Housing and Urban Affairs, Government of India.",
        "MovIqSense SIH 2026 Idea Presentation and Prototype Demonstration, Team MetaMinds 2.0.",
    ]
    for reference in refs:
        add_bullet(doc, reference)

    doc.core_properties.title = "MovIqSense SIH 2026 Project Brief"
    doc.core_properties.subject = "Smart India Hackathon 2026 Problem Statement 26124"
    doc.core_properties.author = "MetaMinds 2.0"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    make_document()
