from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docs" / "MovIqSense_TAM_SAM_SOM_Market_Report.docx"
NAVY, BLUE, LIGHT, TEXT, BORDER = "132238", "1F4E79", "F5F9FC", "1F2937", "D9E2F3"


def cell_shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def cell_format(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), BORDER)
        borders.append(edge)
    props.append(borders)
    margin = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:w"), "110")
        edge.set(qn("w:type"), "dxa")
        margin.append(edge)
    props.append(margin)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def text(p, value, bold=False, color=TEXT, size=10.2):
    run = p.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def body(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.13
    text(p, value)
    return p


def heading(doc, value, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(11 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(value)
    run.bold = True
    run.font.name = "Aptos Display"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.size = Pt(14 if level == 1 else 11.5)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    for i, label in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i])
        cell_shade(c, NAVY)
        cell_format(c)
        p = c.paragraphs[0]
        text(p, label, bold=True, color="FFFFFF", size=9.1)
    for n, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            c = cells[i]
            c.width = Inches(widths[i])
            cell_shade(c, "FFFFFF" if n % 2 == 0 else LIGHT)
            cell_format(c)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            text(p, value, size=8.9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def create():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.6)
    sec.left_margin, sec.right_margin = Inches(0.68), Inches(0.68)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.2)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(fp, "MovIqSense | SIH 2026 | PS 26124 | Illustrative commercial planning model", color="667085", size=8.3)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MovIqSense TAM SAM SOM Market Report")
    r.bold = True
    r.font.name = "Aptos Display"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0, 0, 0)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(sub, "SIH 2026 Problem Statement 26124 | Bharat Electronics Limited | Smart Automation", bold=True, color=BLUE, size=11.5)

    heading(doc, "Purpose and Market Definition")
    body(doc, "This report presents a conservative bottom-up market estimate for MovIqSense, an Edge-AI and GIS platform that turns public buses into mobile road-safety, traffic-intelligence and civic-maintenance sensors. The figures are planning assumptions for SIH evaluation; they are not audited market-size claims. They should be refined with live procurement, fleet and tender data before commercial bidding.")
    body(doc, "The immediate commercial buyer is a government mobility authority: a State Transport Undertaking, Smart City Integrated Command and Control Centre, municipal corporation, Public Works Department or traffic-police authority. The unit of value is a connected bus plus a government dashboard and integration package.")

    heading(doc, "Market Signals")
    table(doc, ["Signal", "Why It Matters to MovIqSense"], [
        ["PM-eBus Sewa targets 10,000 e-buses under a PPP model.", "It creates a visible near-term pool of connected city buses and supports technology-led public transport."],
        ["The scheme includes Intelligent Transit Management System support.", "MovIqSense complements existing mobility infrastructure by adding road, safety and civic intelligence."],
        ["Government focus is moving toward digital urban service delivery.", "A GIS dashboard and structured event data suit command-centre and Smart City workflows."],
    ], [2.85, 4.9])

    heading(doc, "Revenue Assumptions")
    table(doc, ["Commercial Component", "Planning Assumption", "Rationale"], [
        ["Per-bus annual subscription", "₹60,000 per bus per year", "Includes AI event processing, fleet monitoring, alerts and standard dashboard access."],
        ["Authority analytics licence", "₹20-50 lakh per authority per year", "Covers GIS views, priority workflow, reports, role-based access and analytics."],
        ["Initial implementation", "₹20-30 lakh per city or authority", "Covers configuration, integration, training, pilot support and data onboarding."],
        ["Annual maintenance", "15-20% of implementation value", "Supports model updates, security, uptime monitoring and user support."],
    ], [1.9, 2.15, 3.7])

    heading(doc, "TAM SAM SOM Summary")
    table(doc, ["Market", "Definition", "Annual Recurring Revenue Estimate", "Initial Service Opportunity"], [
        ["TAM", "50,000 high-potential public and city-bus units across major STUs, municipal fleets and Smart City corridors.", "₹320 crore", "Additional ₹150 crore potential for multi-city integration and onboarding."],
        ["SAM", "10,000 city e-buses and associated high-readiness authorities in PM-eBus and comparable digital-mobility deployments.", "₹65 crore", "Approx. ₹7.5 crore for first-wave implementation services."],
        ["SOM", "500 buses across 3 early-adopter city or STU deployments within 36 months.", "₹3.6 crore", "Approx. ₹0.6 crore initial implementation services."],
    ], [0.6, 2.9, 2.0, 2.1])

    heading(doc, "Total Addressable Market TAM")
    body(doc, "TAM represents the annual recurring revenue opportunity if MovIqSense is adopted across an estimated 50,000 high-potential public, city and contracted bus units in large urban corridors. The model uses ₹60,000 per bus per year for the core platform. This produces ₹300 crore in annual bus-subscription opportunity. Adding an estimated ₹20 crore in annual authority-level GIS, reporting and analytics licences gives a conservative TAM of approximately ₹320 crore in annual recurring revenue.")
    body(doc, "A separate initial implementation opportunity is estimated at approximately ₹150 crore across phased city and authority deployments. It covers integration with camera systems, GPS or AVL data, Smart City command centres, user training, secure evidence storage and government workflows. It is not included in the recurring revenue total.")

    heading(doc, "Serviceable Available Market SAM")
    body(doc, "SAM focuses on the most accessible market: city e-bus programmes, Smart City corridors and transport authorities already investing in digital fleet operations. The PM-eBus Sewa programme’s 10,000 e-bus target provides a credible initial reference pool. At ₹60,000 per bus per year, this creates ₹60 crore in bus-subscription opportunity. A further ₹5 crore is estimated from authority-level analytics licences, giving a SAM of approximately ₹65 crore in annual recurring revenue.")
    body(doc, "This segment is attractive because it already has public transport operations, digital-transformation intent and a need for Intelligent Transit Management System capability. MovIqSense can be positioned as an AI road-safety and civic-intelligence layer that complements, rather than replaces, existing fleet-management and command-centre systems.")

    heading(doc, "Serviceable Obtainable Market SOM")
    body(doc, "SOM is the realistic market MovIqSense can target in its first 36 months after pilot validation: 500 buses across three early-adopter cities or State Transport Undertaking deployments. At ₹60,000 per bus per year, the platform generates ₹3 crore in annual bus-subscription revenue. Three authority analytics licences at ₹20 lakh per year add ₹0.6 crore, creating an achievable annual recurring revenue target of approximately ₹3.6 crore.")
    body(doc, "The model also includes an estimated ₹0.6 crore in one-time implementation revenue from three ₹20 lakh onboarding and integration engagements. This results in an initial first-year contract opportunity of approximately ₹4.2 crore, while maintaining a recurring-revenue base for long-term support.")

    heading(doc, "Go To Market Strategy")
    table(doc, ["Stage", "Target Customer", "Offer", "Commercial Objective"], [
        ["SIH and prototype", "BEL, Smart City teams, STUs", "Live proof of camera-to-GIS-to-action workflow", "Secure pilot discussions and validation partners."],
        ["Pilot", "One city or STU, 5-10 buses", "3-6 month controlled deployment", "Validate accuracy, alert latency, user adoption and repair workflow."],
        ["Expansion", "3 early-adopter authorities", "100-500 bus subscription plus central dashboard", "Reach ₹3.6 crore annual recurring SOM target."],
        ["Scale", "State and multi-city fleets", "Enterprise integration with ICCC, AVL and PWD", "Build recurring subscriptions and high-value implementation revenue."],
    ], [1.1, 1.65, 2.75, 2.1])

    heading(doc, "Risk Controls and Commercial Credibility")
    body(doc, "MovIqSense should lead with verified operational outcomes rather than untested claims. The first pilot should measure detection precision, false-alert rate, alert latency, number of verified defects, time to authority acknowledgement, number of duplicate events removed and repair closure time. The system should position number-plate output as authorised evidence support requiring human verification, not automatic enforcement.")
    body(doc, "The commercial proposition is strengthened by a low-bandwidth design: the platform transmits structured JSON events and selected evidence clips instead of raw continuous video. This reduces connectivity cost and makes the solution suitable for a staged edge-AI deployment. The initial business model remains Government B2G SaaS plus implementation and annual maintenance, which is appropriate for public-sector procurement and long-term system support.")

    heading(doc, "Sources and Notes")
    for source in [
        "PM-eBus Sewa portal and Government of India scheme materials: 10,000 e-bus deployment target and Intelligent Transit Management System context.",
        "PIB, Cabinet approval for PM-eBus Sewa, 16 August 2023: 10,000 e-buses across 169 cities under the PPP model.",
        "MovIqSense SIH 2026 presentation: prototype scope, value proposition and initial market framing.",
        "All pricing, market segmentation and revenue figures in this report are internal planning assumptions for SIH presentation and must be validated before procurement submission.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        text(p, source, size=9.2)

    doc.core_properties.title = "MovIqSense TAM SAM SOM Market Report"
    doc.core_properties.author = "MetaMinds 2.0"
    doc.core_properties.subject = "SIH 2026 Problem Statement 26124"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    create()
